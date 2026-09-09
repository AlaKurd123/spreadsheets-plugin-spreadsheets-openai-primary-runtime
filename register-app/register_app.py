from __future__ import annotations

import json
import re
from datetime import datetime
from http import HTTPStatus
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

APP_DIR = Path(__file__).resolve().parent
STATIC_DIR = APP_DIR.parent / "docs"
EXPORT_DIR = APP_DIR / "exports"
def safe_filename(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9_-]+", "-", value.strip()).strip("-")
    return cleaned or "register"


def make_workbook(data: dict) -> Path:
    EXPORT_DIR.mkdir(exist_ok=True)
    now = datetime.now()
    class_name = str(data["className"]).strip()
    date = str(data["date"]).strip()
    teacher = str(data["teacher"]).strip()
    rows = data["students"]

    book = Workbook()
    sheet = book.active
    sheet.title = "Register"
    sheet.merge_cells("A1:D1")
    sheet["A1"] = "Temporary Register"
    sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    sheet["A1"].alignment = Alignment(horizontal="center")

    details = [("Class", class_name), ("Date", date), ("Teacher", teacher), ("Submitted", now.strftime("%d %b %Y %H:%M"))]
    for row_number, (label, value) in enumerate(details, start=3):
        sheet.cell(row=row_number, column=1, value=label).font = Font(bold=True)
        sheet.cell(row=row_number, column=2, value=value)

    headers = ["Student name", "Status", "Time", "Notes"]
    header_row = 8
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B9BD5")
        cell.alignment = Alignment(horizontal="center")

    for row_number, student in enumerate(rows, start=header_row + 1):
        values = [student["name"], student["status"], student.get("time", ""), student.get("notes", "")]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_number, column=column, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if student["status"] == "Absent":
            sheet.cell(row=row_number, column=2).fill = PatternFill("solid", fgColor="FCE4D6")
        elif student["status"] == "Late":
            sheet.cell(row=row_number, column=2).fill = PatternFill("solid", fgColor="FFF2CC")

    sheet.freeze_panes = "A9"
    widths = [28, 14, 14, 45]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    filename = f"{safe_filename(class_name)}_{date}_{now.strftime('%H%M%S')}.xlsx"
    output = EXPORT_DIR / filename
    book.save(output)
    return output


class RegisterHandler(SimpleHTTPRequestHandler):
    def do_GET(self):
        requested = urlparse(self.path).path
        if requested in ("/", "/index.html"):
            self.path = "/index.html"
            return self.serve_static()
        if requested == "/health":
            return self.respond_json(HTTPStatus.OK, {"ok": True})
        if requested in ("/app.js", "/styles.css"):
            self.path = requested
            return self.serve_static()
        self.send_error(HTTPStatus.NOT_FOUND, "Not found")

    def do_POST(self):
        if urlparse(self.path).path != "/submit":
            self.send_error(HTTPStatus.NOT_FOUND, "Not found")
            return
        try:
            content_length = int(self.headers.get("Content-Length", "0"))
            payload = json.loads(self.rfile.read(content_length))
            data = self.validate(payload)
            workbook = make_workbook(data)
            self.respond_workbook(workbook)
        except (ValueError, KeyError, json.JSONDecodeError) as error:
            self.respond_json(HTTPStatus.BAD_REQUEST, {"error": str(error)})
        except Exception as error:
            self.respond_json(HTTPStatus.INTERNAL_SERVER_ERROR, {"error": f"Could not export the register: {error}"})

    def validate(self, payload: dict) -> dict:
        if not isinstance(payload, dict):
            raise ValueError("Invalid form submission.")
        data = {key: str(payload.get(key, "")).strip() for key in ("className", "date", "teacher")}
        if not all(data.values()):
            raise ValueError("Class, date, and teacher are required.")
        students = payload.get("students")
        if not isinstance(students, list) or not students:
            raise ValueError("Add at least one student.")
        cleaned_students = []
        for student in students:
            name = str(student.get("name", "")).strip()
            status = str(student.get("status", "Present")).strip()
            if not name:
                continue
            if status not in ("Present", "Absent", "Late"):
                raise ValueError("Each student must have a valid status.")
            cleaned_students.append({
                "name": name,
                "status": status,
                "time": str(student.get("time", "")).strip(),
                "notes": str(student.get("notes", "")).strip(),
            })
        if not cleaned_students:
            raise ValueError("Add at least one student name.")
        data["students"] = cleaned_students
        return data

    def serve_static(self):
        target = STATIC_DIR / self.path.lstrip("/")
        if not target.is_file() or STATIC_DIR not in target.resolve().parents:
            self.send_error(HTTPStatus.NOT_FOUND, "Not found")
            return
        content_type = "text/html; charset=utf-8" if target.suffix == ".html" else "text/css; charset=utf-8" if target.suffix == ".css" else "application/javascript; charset=utf-8"
        body = target.read_bytes()
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def respond_json(self, status: HTTPStatus, value: dict):
        body = json.dumps(value).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def respond_workbook(self, workbook: Path):
        body = workbook.read_bytes()
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.send_header("Content-Disposition", f'attachment; filename="{workbook.name}"')
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)


if __name__ == "__main__":
    print("Temporary Register is running at http://127.0.0.1:8081")
    ThreadingHTTPServer(("127.0.0.1", 8081), RegisterHandler).serve_forever()
